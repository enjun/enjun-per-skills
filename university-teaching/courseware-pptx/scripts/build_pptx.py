# -*- coding: utf-8 -*-
"""
courseware-pptx 通用引擎：从 JSON 课件数据生成 16:9 教学 PPT。

用法:
    uv run --with python-pptx python build_pptx.py <content.json> [output.pptx]

JSON schema 见 references/schema.md，完整示例见 examples/week1_python.json。
"""
import json
import os
import sys
from dataclasses import dataclass, replace

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# ============== 页面尺寸（16:9） ==============
PAGE_W = 13.333
PAGE_H = 7.5

# ============== 主题色 ==============
PY_BLUE = RGBColor(0x30, 0x69, 0x98)
PY_BLUE_LIGHT = RGBColor(0x4B, 0x8B, 0xC8)
PY_YELLOW = RGBColor(0xFF, 0xD4, 0x3B)
PY_DARK = RGBColor(0x1E, 0x2A, 0x38)
PY_GRAY = RGBColor(0x5A, 0x66, 0x73)
PY_BG_LIGHT = RGBColor(0xF5, 0xF7, 0xFA)
PY_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PY_CODE_BG = RGBColor(0x2B, 0x3A, 0x45)
PY_CODE_FG = RGBColor(0xE8, 0xE8, 0xE8)
PY_ACCENT_GREEN = RGBColor(0x37, 0xB2, 0x4D)
PY_ACCENT_RED = RGBColor(0xE6, 0x4A, 0x4A)
PY_ACCENT_ORANGE = RGBColor(0xF5, 0x9E, 0x42)
PY_PURPLE = RGBColor(0x9B, 0x59, 0xB6)

# 引擎 chrome 固定色（不出现在 JSON 中）
CARD_BORDER = RGBColor(0xDD, 0xE3, 0xEA)
ROW_BORDER = RGBColor(0xE5, 0xE9, 0xEF)
AXIS_GRAY = RGBColor(0xD8, 0xDF, 0xE6)
CODE_TAB_FG = RGBColor(0xB8, 0xC2, 0xCC)
COVER_SUBTLE = RGBColor(0xC9, 0xD8, 0xE6)
DOT_RED = RGBColor(0xFF, 0x5F, 0x56)
DOT_YELLOW = RGBColor(0xFF, 0xBD, 0x2E)
DOT_GREEN = RGBColor(0x27, 0xC9, 0x3F)

CN_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"


class Theme:
    """palette 名 → RGBColor；JSON 中颜色只允许 palette 名。"""

    NAMES = {
        "blue": PY_BLUE,
        "blue_light": PY_BLUE_LIGHT,
        "yellow": PY_YELLOW,
        "dark": PY_DARK,
        "gray": PY_GRAY,
        "bg_light": PY_BG_LIGHT,
        "white": PY_WHITE,
        "code_bg": PY_CODE_BG,
        "code_fg": PY_CODE_FG,
        "green": PY_ACCENT_GREEN,
        "red": PY_ACCENT_RED,
        "orange": PY_ACCENT_ORANGE,
        "purple": PY_PURPLE,
    }

    def __init__(self, overrides=None):
        self._map = dict(self.NAMES)
        for name, hexstr in (overrides or {}).items():
            if name not in self._map:
                raise ValueError(f"palette 覆盖使用了未知色名 '{name}'，合法名: {sorted(self._map)}")
            self._map[name] = RGBColor.from_string(hexstr.lstrip("#").upper())

    def c(self, name, fallback=None):
        key = name if name is not None else fallback
        if key is None:
            return None
        if key not in self._map:
            raise ValueError(f"未知颜色名 '{key}'，合法名: {sorted(self._map)}")
        return self._map[key]


@dataclass
class Ctx:
    page_no: int
    total: int
    course_line: str
    theme: Theme


# ============== 基础助手 ==============

def _apply_font(run, font_name, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font_name)


def _add_rich_runs(p, text, font_name, size, color, bold=False):
    """按 **bold** 迷你标记拆分文本并添加 run。"""
    parts = text.split("**")
    if len(parts) == 1:
        run = p.add_run()
        run.text = text
        _apply_font(run, font_name, size, color, bold)
        return
    for j, part in enumerate(parts):
        if not part:
            continue
        run = p.add_run()
        run.text = part
        _apply_font(run, font_name, size, color, bold or (j % 2 == 1))


def add_rect(slide, x, y, w, h, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.75)
    try:
        shape.shadow.inherit = False
    except Exception:
        pass
    return shape


def add_round_rect(slide, x, y, w, h, fill_color, line_color=None, corner=0.08):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.75)
    shape.adjustments[0] = corner
    try:
        shape.shadow.inherit = False
    except Exception:
        pass
    return shape


def add_text(slide, x, y, w, h, text, font_size=18, bold=False, color=PY_DARK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font_name=CN_FONT,
             line_spacing=1.2):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        _apply_font(run, font_name, font_size, color, bold)
    return tb


def add_bullets(slide, x, y, w, h, items, font_size=16, color=PY_DARK,
                line_spacing=1.25, bullet_color=PY_BLUE_LIGHT, font_name=CN_FONT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        bullet_run = p.add_run()
        bullet_run.text = "●  "
        _apply_font(bullet_run, font_name, font_size, bullet_color, bold=True)
        if isinstance(item, dict):
            main_text, sub_text = item.get("main", ""), item.get("sub", "")
            _add_rich_runs(p, main_text, font_name, font_size, color, bold=True)
            if sub_text:
                _add_rich_runs(p, sub_text, font_name, font_size - 1, PY_GRAY)
        else:
            _add_rich_runs(p, item, font_name, font_size, color)
    return tb


def add_code_block(slide, x, y, w, h, code, font_size=12, label="python"):
    title_h = 0.32
    add_round_rect(slide, x, y, w, title_h, PY_CODE_BG, corner=0.5)
    for i, c in enumerate([DOT_RED, DOT_YELLOW, DOT_GREEN]):
        d = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x + 0.18 + 0.18 * i), Inches(y + 0.10),
            Inches(0.13), Inches(0.13))
        d.fill.solid()
        d.fill.fore_color.rgb = c
        d.line.fill.background()
    add_text(slide, x + 0.85, y, w - 0.9, title_h, label,
             font_size=11, color=CODE_TAB_FG, anchor=MSO_ANCHOR.MIDDLE)
    code_y = y + title_h
    code_h = h - title_h
    add_rect(slide, x, code_y, w, code_h, PY_CODE_BG)
    tb = slide.shapes.add_textbox(Inches(x + 0.18), Inches(code_y + 0.08),
                                  Inches(w - 0.3), Inches(code_h - 0.15))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(code.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.15
        r = p.add_run()
        r.text = line if line else " "
        _apply_font(r, CODE_FONT, font_size, PY_CODE_FG)
    return tb


def add_page_header(slide, ctx, section_title=""):
    add_rect(slide, 0, 0, 0.25, PAGE_H, ctx.theme.c("blue"))
    add_text(slide, PAGE_W - 1.4, 0.18, 1.2, 0.3,
             f"{ctx.page_no:02d} / {ctx.total:02d}", font_size=10,
             color=ctx.theme.c("gray"), align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 0.45, 0.15, 6, 0.3, ctx.course_line, font_size=10,
             color=ctx.theme.c("gray"), anchor=MSO_ANCHOR.MIDDLE)
    if section_title:
        add_text(slide, 6.5, 0.15, 5.5, 0.3, section_title, font_size=10,
                 color=ctx.theme.c("blue_light"), align=PP_ALIGN.RIGHT,
                 anchor=MSO_ANCHOR.MIDDLE, bold=True)


def add_section_title(slide, theme, title, subtitle="", y_offset=0.7):
    add_rect(slide, 0.6, y_offset + 0.05, 0.12, 0.7, theme.c("yellow"))
    add_text(slide, 0.85, y_offset, 11.5, 0.8, title, font_size=32, bold=True,
             color=theme.c("blue"), anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, 0.85, y_offset + 0.85, 11.5, 0.4, subtitle,
                 font_size=14, color=theme.c("gray"))


def add_footer(slide, theme):
    add_rect(slide, 0, PAGE_H - 0.04, PAGE_W, 0.04, theme.c("blue"))


def content_slide(prs, spec, ctx):
    """内容页工厂：白底 → 页眉 → 页脚 → 大标题。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, PAGE_W, PAGE_H, ctx.theme.c("white"))
    add_page_header(slide, ctx, spec.get("section", ""))
    add_footer(slide, ctx.theme)
    if spec.get("title"):
        add_section_title(slide, ctx.theme, spec["title"], spec.get("subtitle", ""))
    return slide


# ============== 卡片与内容块 ==============

def _est_height(block):
    """未显式给出 height 时的估算高度（英寸）。"""
    kind = block["kind"]
    if kind == "bullets":
        n = len(block["items"])
        size = block.get("size", 15)
        ls = block.get("line_spacing", 1.45)
        return round(n * size / 72.0 * ls * 1.75 + 0.1, 2)
    if kind == "code":
        n = block["code"].count("\n") + 1
        size = block.get("size", 12)
        return round(0.32 + n * size / 72.0 * 1.15 * 1.5 + 0.15, 2)
    if kind == "text":
        n = block["text"].count("\n") + 1
        size = block.get("size", 12)
        return round(n * size / 72.0 * 1.4 + 0.08, 2)
    if kind == "table":
        return len(block["rows"]) * block.get("row_h", 0.34)
    if kind == "callout":
        return 0.85
    if kind == "kv_code":
        return len(block["rows"]) * block.get("row_h", 0.42) + 0.05
    if kind == "tagged_items":
        n = len(block["items"])
        return n * block.get("item_h", 1.25) + (n - 1) * 0.1
    raise ValueError(f"未知内容块类型 '{kind}'")


def render_block(slide, block, x, y, w, theme):
    """渲染一个内容块，返回排流后的新 y。坐标均为英寸浮点数。"""
    kind = block["kind"]
    h = block.get("height") or _est_height(block)
    if kind == "bullets":
        add_bullets(slide, x, y, w, h, block["items"],
                    font_size=block.get("size", 15),
                    line_spacing=block.get("line_spacing", 1.45),
                    bullet_color=theme.c(block.get("bullet_color"), fallback="blue_light"))
    elif kind == "code":
        add_code_block(slide, x, y, w, h, block["code"],
                       font_size=block.get("size", 12),
                       label=block.get("label", "python"))
    elif kind == "text":
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.05)
        tf.margin_right = Inches(0.05)
        tf.margin_top = Inches(0.02)
        tf.margin_bottom = Inches(0.02)
        for i, line in enumerate(block["text"].split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = 1.3
            _add_rich_runs(p, line, CN_FONT, block.get("size", 12),
                           theme.c(block.get("color"), fallback="dark"),
                           block.get("bold", False))
    elif kind == "table":
        widths = block["widths"]
        row_h = block.get("row_h", 0.34)
        size = block.get("size", 11)
        col_styles = block.get("col_styles") or []
        for i, row in enumerate(block["rows"]):
            ry = y + row_h * i
            bg = theme.c("bg_light") if i % 2 == 0 else theme.c("white")
            add_rect(slide, x, ry, sum(widths), row_h, bg, line_color=ROW_BORDER)
            cx = x + 0.05
            for j, cell in enumerate(row):
                cw = widths[j]
                style = col_styles[j] if j < len(col_styles) and col_styles[j] else {}
                if isinstance(cell, dict):
                    text = cell["text"]
                    color = theme.c(cell.get("color"), fallback=style.get("color", "dark"))
                    bold = cell.get("bold", style.get("bold", False))
                else:
                    text = cell
                    color = theme.c(style.get("color", "dark"))
                    bold = style.get("bold", False)
                add_text(slide, cx, ry, cw, row_h, text, font_size=size,
                         color=color, bold=bold, anchor=MSO_ANCHOR.MIDDLE)
                cx += cw
    elif kind == "callout":
        add_round_rect(slide, x, y, w, h, theme.c("yellow"), corner=0.2)
        add_text(slide, x, y, w, h, block["text"], font_size=12, bold=True,
                 color=theme.c("blue"), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    elif kind == "kv_code":
        row_h = block.get("row_h", 0.42)
        for i, (label, command) in enumerate(block["rows"]):
            ry = y + row_h * i
            bg = theme.c("bg_light") if i % 2 == 0 else theme.c("white")
            add_rect(slide, x, ry, w, row_h, bg, line_color=ROW_BORDER)
            add_text(slide, x + 0.1, ry, 2.3, row_h, label, font_size=12,
                     color=theme.c("dark"), anchor=MSO_ANCHOR.MIDDLE)
            add_round_rect(slide, x + 2.45, ry + 0.06, w - 2.65, row_h - 0.12,
                           theme.c("code_bg"), corner=0.25)
            add_text(slide, x + 2.55, ry + 0.06, w - 2.85, row_h - 0.12, command,
                     font_size=11, color=theme.c("code_fg"), font_name=CODE_FONT,
                     anchor=MSO_ANCHOR.MIDDLE)
    elif kind == "tagged_items":
        item_h = block.get("item_h", 1.25)
        tag_color = theme.c(block.get("tag_color"), fallback="green")
        for i, item in enumerate(block["items"]):
            ry = y + (item_h + 0.1) * i
            add_round_rect(slide, x, ry + 0.15, 0.65, 0.4, tag_color, corner=0.3)
            add_text(slide, x, ry + 0.15, 0.65, 0.4, item["tag"], font_size=10,
                     bold=True, color=theme.c("white"), align=PP_ALIGN.CENTER,
                     anchor=MSO_ANCHOR.MIDDLE)
            add_text(slide, x + 0.75, ry + 0.1, w - 0.9, 0.4, item["title"],
                     font_size=13, bold=True, color=theme.c("dark"))
            add_text(slide, x + 0.75, ry + 0.5, w - 0.9, 0.75, item["text"],
                     font_size=11, color=theme.c("gray"), line_spacing=1.3)
    else:
        raise ValueError(f"未知内容块类型 '{kind}'")
    return y + h + 0.15


def render_card(slide, card, x, y, w, h, theme):
    """渲染一张卡片（含标题 chrome）并排流其内容块。"""
    style = card.get("style", "plain")
    blocks = card.get("blocks", [])
    corner = 0.04
    if style == "header":
        header_h = card.get("header_h", 0.7)
        accent = theme.c(card.get("accent"), fallback="blue")
        title_color = theme.c(card.get("title_color"), fallback="white")
        add_round_rect(slide, x, y, w, h, theme.c("white"), line_color=CARD_BORDER, corner=corner)
        add_rect(slide, x, y, w, header_h, accent)
        title_size = card.get("title_size", 17 if header_h >= 0.7 else 15)
        badge = card.get("badge")
        title_w = w - 0.6 - (1.0 if badge else 0.0)
        add_text(slide, x + 0.3, y, title_w, header_h, card["title"],
                 font_size=title_size, bold=True, color=title_color,
                 anchor=MSO_ANCHOR.MIDDLE)
        if badge:
            bx = x + w - 0.7
            add_round_rect(slide, bx, y + 0.13, 0.44, 0.44,
                           theme.c(badge.get("bg", "yellow")), corner=0.5)
            add_text(slide, bx, y + 0.13, 0.44, 0.44, badge["text"],
                     font_size=16 if len(badge["text"]) <= 1 else 11,
                     bold=True, color=theme.c(badge.get("fg", "blue")),
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        by = y + header_h + 0.25
    elif style == "strip":
        accent = theme.c(card.get("accent"), fallback="blue")
        add_round_rect(slide, x, y, w, h, theme.c("white"), line_color=CARD_BORDER, corner=corner)
        add_rect(slide, x, y, 0.12, h, accent)
        add_text(slide, x + 0.35, y + 0.15, w - 0.7, 0.5, card["title"],
                 font_size=card.get("title_size", 15), bold=True,
                 color=theme.c(card.get("title_color"), fallback="blue"))
        by = y + 0.85
    else:  # plain
        add_round_rect(slide, x, y, w, h, theme.c(card.get("bg"), fallback="bg_light"), corner=corner)
        add_text(slide, x + 0.3, y + 0.2, w - 0.6, 0.5, card["title"],
                 font_size=card.get("title_size", 14), bold=True,
                 color=theme.c(card.get("title_color"), fallback="blue"))
        by = y + 0.85
    for block in blocks:
        by = render_block(slide, block, x + 0.3, by, w - 0.6, theme)


# ============== 17 种版式 builder ==============

def build_cover(prs, spec, ctx):
    theme = ctx.theme
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, PAGE_W, PAGE_H, theme.c("blue"))
    for tx, tw, color in [(8.5, 4.83, "blue_light"), (10.5, 2.83, "yellow")]:
        tri = slide.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(tx), Inches(0),
                                     Inches(tw), Inches(PAGE_H))
        tri.fill.solid()
        tri.fill.fore_color.rgb = theme.c(color)
        tri.line.fill.background()
        try:
            tri.shadow.inherit = False
        except Exception:
            pass
    for cx, cy, sz in [(1.2, 1.2, 0.4), (1.8, 0.7, 0.18), (2.3, 1.5, 0.22)]:
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx), Inches(cy), Inches(sz), Inches(sz))
        c.fill.solid()
        c.fill.fore_color.rgb = theme.c("yellow")
        c.line.fill.background()
    add_text(slide, 0.8, 2.0, 3, 0.6, spec["kicker"], font_size=20, bold=True,
             color=theme.c("yellow"))
    for i, line in enumerate(spec["title_lines"]):
        add_text(slide, 0.8, 2.6 + 1.1 * i, 9.5, 1.4, line, font_size=54,
                 bold=True, color=theme.c("white"))
    add_rect(slide, 0.8, 5.05, 1.2, 0.06, theme.c("yellow"))
    add_text(slide, 0.8, 5.2, 9.5, 0.5, spec["subtitle"], font_size=18,
             color=theme.c("yellow"))
    add_text(slide, 0.8, 5.8, 9.5, 0.5,
             spec.get("course_line", ctx.course_line.replace("  ·  ", " · ")),
             font_size=16, color=COVER_SUBTLE)
    add_text(slide, 0.8, 6.8, 9, 0.4, spec["credit"], font_size=12,
             color=theme.c("yellow"))


def build_divider(prs, spec, ctx):
    theme = ctx.theme
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, PAGE_W, PAGE_H, theme.c("white"))
    add_rect(slide, 0, 0, 5.5, PAGE_H, theme.c(spec["color"]))
    add_text(slide, 0.6, 1.6, 5, 0.6, "PART", font_size=18, bold=True,
             color=theme.c("yellow"))
    add_text(slide, 0.6, 2.4, 5, 2.0, spec["no"], font_size=160, bold=True,
             color=theme.c("white"))
    add_text(slide, PAGE_W - 1.4, 0.18, 1.2, 0.3,
             f"{ctx.page_no:02d} / {ctx.total:02d}", font_size=10,
             color=theme.c("gray"), align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, 6.3, 0.15, 6, 0.3, ctx.course_line, font_size=10,
             color=theme.c("gray"), anchor=MSO_ANCHOR.MIDDLE)
    add_rect(slide, 6.0, 3.4, 0.12, 0.7, theme.c("yellow"))
    add_text(slide, 6.3, 3.3, 6.5, 0.9, spec["title"], font_size=36, bold=True,
             color=theme.c("dark"))
    add_text(slide, 6.3, 4.15, 6.5, 0.6, spec["subtitle"], font_size=16,
             color=theme.c("gray"))


def build_agenda(prs, spec, ctx):
    theme = ctx.theme
    slide = content_slide(prs, spec, ctx)
    items = spec["items"]
    add_rect(slide, 1.5, 2.0, 0.04, 0.85 * (len(items) - 1) + 0.5, AXIS_GRAY)
    for i, item in enumerate(items):
        y = 2.0 + 0.85 * i
        color = theme.c(item.get("color"), fallback="blue")
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.3), Inches(y + 0.05),
                                   Inches(0.44), Inches(0.44))
        c.fill.solid()
        c.fill.fore_color.rgb = color
        c.line.color.rgb = theme.c("white")
        c.line.width = Pt(2)
        add_text(slide, 1.3, y + 0.05, 0.44, 0.44, item["no"], font_size=11,
                 bold=True, color=theme.c("white"), align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, 2.1, y, 4.5, 0.45, item["title"], font_size=20,
                 bold=True, color=theme.c("blue"), anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, 2.1, y + 0.42, 10, 0.35, item["desc"], font_size=13,
                 color=theme.c("gray"))
    sidebar = spec["sidebar"]
    add_round_rect(slide, 10.0, 2.0, 2.85, 4.2, theme.c("blue"), corner=0.05)
    add_text(slide, 10.25, 2.3, 2.4, 0.4, sidebar["heading"], font_size=14,
             bold=True, color=theme.c("yellow"))
    steps_text = "\n".join(
        [sidebar["steps"][0]] + [f"    ↓\n{s}" for s in sidebar["steps"][1:]])
    add_text(slide, 10.25, 2.85, 2.4, 3.0, steps_text, font_size=13,
             color=theme.c("white"), line_spacing=1.4)


def build_two_cards(prs, spec, ctx):
    slide = content_slide(prs, spec, ctx)
    cards = spec["cards"]
    widths = spec.get("widths", [6.0, 5.9])
    y = 1.85
    h = spec.get("height", 4.7)
    render_card(slide, cards[0], 0.6, y, widths[0], h, ctx.theme)
    render_card(slide, cards[1], PAGE_W - 0.6 - widths[1], y, widths[1], h, ctx.theme)


def build_story_person(prs, spec, ctx):
    theme = ctx.theme
    slide = content_slide(prs, spec, ctx)
    y, h = 1.85, 4.7
    story = spec["story"]
    x, w = 0.6, 7.2
    add_round_rect(slide, x, y, w, h, theme.c("white"), line_color=CARD_BORDER, corner=0.04)
    add_rect(slide, x, y, w, 0.08, theme.c("yellow"))
    add_text(slide, x + 0.4, y + 0.3, w - 0.8, 0.9, story["headline"],
             font_size=42, bold=True, color=theme.c("blue"))
    add_text(slide, x + 0.4, y + 1.25, w - 0.8, 0.5, story["lead"],
             font_size=18, color=theme.c("dark"))
    add_bullets(slide, x + 0.4, y + 1.85, w - 0.8, h - 2.0, story["bullets"],
                font_size=13, line_spacing=1.45,
                bullet_color=theme.c(story.get("bullet_color"), fallback="yellow"))
    person = spec["person"]
    x2, w2 = 8.05, 4.7
    add_round_rect(slide, x2, y, w2, h, theme.c("blue"), corner=0.04)
    av = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x2 + w2 / 2 - 0.9),
                                Inches(y + 0.4), Inches(1.8), Inches(1.8))
    av.fill.solid()
    av.fill.fore_color.rgb = theme.c("yellow")
    av.line.color.rgb = theme.c("white")
    av.line.width = Pt(3)
    add_text(slide, x2 + w2 / 2 - 0.9, y + 0.4, 1.8, 1.8, person["avatar"],
             font_size=42, bold=True, color=theme.c("blue"),
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, x2 + 0.3, y + 2.45, w2 - 0.6, 0.5, person["name"],
             font_size=20, bold=True, color=theme.c("white"), align=PP_ALIGN.CENTER)
    add_text(slide, x2 + 0.3, y + 2.9, w2 - 0.6, 0.4, person["role"],
             font_size=12, color=theme.c("yellow"), align=PP_ALIGN.CENTER)
    add_round_rect(slide, x2 + 0.3, y + 3.45, w2 - 0.6, 1.05, theme.c("white"), corner=0.15)
    add_text(slide, x2 + 0.45, y + 3.55, w2 - 0.9, 0.95, person["quote"],
             font_size=12, color=theme.c("dark"), align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)


def _rich_cell(cell, theme, default_style=None):
    default_style = default_style or {}
    if isinstance(cell, dict):
        return (cell["text"],
                theme.c(cell.get("color"), fallback=default_style.get("color", "dark")),
                cell.get("bold", default_style.get("bold", False)))
    return (cell, theme.c(default_style.get("color", "dark")),
            default_style.get("bold", False))


def build_table(prs, spec, ctx):
    theme = ctx.theme
    slide = content_slide(prs, spec, ctx)
    widths = spec["widths"]
    aligns = spec.get("aligns") or ["center"] * len(widths)
    row_h = spec.get("row_h", 0.55)
    header_h = spec.get("header_h", 0.5)
    col_styles = spec.get("col_styles") or []
    first_col_bold = spec.get("first_col_bold", False)
    x, total_w = 0.6, sum(widths)
    header_y = 1.85
    add_rect(slide, x, header_y, total_w, header_h, theme.c("blue"))
    cx = x
    for j, hcell in enumerate(spec["headers"]):
        text, color, bold = _rich_cell(hcell, theme)
        align = PP_ALIGN.LEFT if aligns[j] == "left" else PP_ALIGN.CENTER
        add_text(slide, cx, header_y, widths[j], header_h, text, font_size=15,
                 bold=True, color=color, align=align, anchor=MSO_ANCHOR.MIDDLE)
        cx += widths[j]
    for i, row in enumerate(spec["rows"]):
        ry = header_y + header_h + row_h * i
        bg = theme.c("white") if i % 2 == 0 else theme.c("bg_light")
        add_rect(slide, x, ry, total_w, row_h, bg, line_color=ROW_BORDER)
        cx = x
        for j, cell in enumerate(row):
            style = dict(col_styles[j]) if j < len(col_styles) and col_styles[j] else {}
            if j == 0 and first_col_bold:
                style.setdefault("bold", True)
            text, color, bold = _rich_cell(cell, theme, style)
            align = PP_ALIGN.LEFT if aligns[j] == "left" else PP_ALIGN.CENTER
            add_text(slide, cx, ry, widths[j], row_h, text, font_size=13,
                     color=color, bold=bold, align=align, anchor=MSO_ANCHOR.MIDDLE)
            cx += widths[j]
    note = spec.get("note")
    if note:
        note_y = min(6.45, header_y + header_h + row_h * len(spec["rows"]) + 0.25)
        add_round_rect(slide, x, note_y, total_w, 0.5, theme.c("bg_light"),
                       line_color=theme.c("yellow"), corner=0.3)
        tb = slide.shapes.add_textbox(Inches(x + 0.25), Inches(note_y),
                                      Inches(total_w - 0.45), Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        _add_rich_runs(p, note, CN_FONT, 13, theme.c("blue"), bold=True)


def build_stats(prs, spec, ctx):
    theme = ctx.theme
    slide = content_slide(prs, spec, ctx)
    cards = spec["cards"]
    card_w, card_h, gap = 4.0, 2.4, 0.2
    total_w = card_w * len(cards) + gap * (len(cards) - 1)
    start_x = (PAGE_W - total_w) / 2
    y = 2.0
    for i, card in enumerate(cards):
        cx = start_x + (card_w + gap) * i
        color = theme.c(card.get("color"), fallback="blue")
        add_round_rect(slide, cx, y, card_w, card_h, theme.c("white"),
                       line_color=CARD_BORDER, corner=0.05)
        add_rect(slide, cx, y, card_w, 0.12, color)
        add_text(slide, cx, y + 0.35, card_w, 0.9, card["icon"], font_size=44,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, cx, y + 1.3, card_w, 0.5, card["title"], font_size=20,
                 bold=True, color=color, align=PP_ALIGN.CENTER)
        add_text(slide, cx + 0.3, y + 1.85, card_w - 0.6, 0.5, card["desc"],
                 font_size=12, color=theme.c("gray"), align=PP_ALIGN.CENTER)
    banner = spec["banner"]
    y2 = 4.7
    add_round_rect(slide, 0.6, y2, 12.15, 2.0, theme.c("blue"), corner=0.04)
    add_text(slide, 0.6, y2 + 0.15, 12.15, 0.5, banner["heading"], font_size=18,
             bold=True, color=theme.c("yellow"), align=PP_ALIGN.CENTER)
    items = banner["items"]
    item_w, item_gap = 2.85, 0.1
    item_total = item_w * len(items) + item_gap * (len(items) - 1)
    item_start = (PAGE_W - item_total) / 2
    iy = y2 + 0.85
    for i, item in enumerate(items):
        ix = item_start + (item_w + item_gap) * i
        add_text(slide, ix, iy, item_w, 0.5, item["icon"], font_size=24,
                 align=PP_ALIGN.CENTER)
        add_text(slide, ix, iy + 0.55, item_w, 0.4, item["title"], font_size=14,
                 bold=True, color=theme.c("yellow"), align=PP_ALIGN.CENTER)
        add_text(slide, ix, iy + 0.95, item_w, 0.4, item["desc"], font_size=10,
                 color=theme.c("white"), align=PP_ALIGN.CENTER)


def build_ai_prompts(prs, spec, ctx):
    theme = ctx.theme
    slide = content_slide(prs, spec, ctx)
    add_round_rect(slide, 0.6, 1.85, 12.15, 4.7, theme.c("white"),
                   line_color=CARD_BORDER, corner=0.04)
    add_rect(slide, 0.6, 1.85, 12.15, 0.7, theme.c("orange"))
    add_text(slide, 0.85, 1.85, 11.9, 0.7,
             spec.get("box_title", "💬  推荐 Prompt"), font_size=16, bold=True,
             color=theme.c("white"), anchor=MSO_ANCHOR.MIDDLE)
    prompts = spec["prompts"]
    by, bh = 2.75, min(0.85, 3.6 / len(prompts) - 0.05)
    for i, prompt in enumerate(prompts):
        ry = by + (bh + 0.05) * i
        add_round_rect(slide, 0.85, ry + 0.1, 1.2, 0.6, theme.c("orange"), corner=0.3)
        add_text(slide, 0.85, ry + 0.1, 1.2, 0.6, prompt["label"], font_size=11,
                 bold=True, color=theme.c("white"), align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, 2.25, ry, 10.3, bh, prompt["text"], font_size=14,
                 color=theme.c("dark"), anchor=MSO_ANCHOR.MIDDLE)
        if i < len(prompts) - 1:
            add_rect(slide, 0.85, ry + bh, 11.6, 0.01, ROW_BORDER)


def build_code_compare(prs, spec, ctx):
    theme = ctx.theme
    slide = content_slide(prs, spec, ctx)
    cards = spec["cards"]
    cw, gap, ch = 4.0, 0.18, 4.5
    total_w = cw * len(cards) + gap * (len(cards) - 1)
    start_x = (PAGE_W - total_w) / 2
    y = 1.85
    for i, card in enumerate(cards):
        cx = start_x + (cw + gap) * i
        color = theme.c(card.get("color"), fallback="blue")
        add_round_rect(slide, cx, y, cw, ch, theme.c("white"),
                       line_color=CARD_BORDER, corner=0.03)
        add_rect(slide, cx, y, cw, 0.45, color)
        add_text(slide, cx, y, cw, 0.45, card["lang"], font_size=16, bold=True,
                 color=theme.c("white"), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_code_block(slide, cx + 0.15, y + 0.6, cw - 0.3, ch - 0.75,
                       card["code"], font_size=11, label=card["filename"])
    notes = spec["notes"]
    for i, note in enumerate(notes):
        lx = 0.6 + 4.18 * i
        color = theme.c(note["color"]) if note.get("color") else theme.c("gray")
        add_text(slide, lx, 6.5, 4.0, 0.4, f"{note.get('icon', '')}  {note['text']}",
                 font_size=12, color=color, anchor=MSO_ANCHOR.MIDDLE)


def build_code_demo(prs, spec, ctx):
    slide = content_slide(prs, spec, ctx)
    render_card(slide, spec["scenario"], 0.6, 1.85, 4.3, 4.7, ctx.theme)
    code = spec["code"]
    add_code_block(slide, 5.1, 1.85, 7.65, 4.7, code["code"],
                   font_size=code.get("size", 10), label=code.get("filename", "demo.py"))


def build_code_full(prs, spec, ctx):
    """整页代码块：用于放不进卡片的长代码。代码过长时应在 JSON 层
    拆成连续多页（filename 标注 （1/2）（2/2）），禁止删改代码本身。"""
    slide = content_slide(prs, spec, ctx)
    code = spec["code"]
    size = code.get("size", 11)
    n = code["code"].count("\n") + 1
    y_start = 2.05 if spec.get("subtitle") else 1.8
    h = code.get("height", 0.47 + n * size * 1.38 / 72 + 0.08)
    h = min(h, 7.25 - y_start)
    add_code_block(slide, 0.6, y_start, 12.15, h, code["code"],
                   font_size=size, label=code.get("filename", "demo.py"))


def build_icon_grid(prs, spec, ctx):
    theme = ctx.theme
    slide = content_slide(prs, spec, ctx)
    items = spec["items"]
    cols = spec.get("cols", 3)
    rows = (len(items) + cols - 1) // cols
    cw, ch = 4.0, 1.45
    gap_x, gap_y = 0.15, 0.12
    grid_w = cw * cols + gap_x * (cols - 1)
    start_x = (PAGE_W - grid_w) / 2
    start_y = 1.85
    for i, item in enumerate(items):
        x = start_x + (i % cols) * (cw + gap_x)
        y = start_y + (i // cols) * (ch + gap_y)
        add_round_rect(slide, x, y, cw, ch, theme.c("white"),
                       line_color=CARD_BORDER, corner=0.08)
        add_rect(slide, x, y, 0.12, ch, theme.c(item.get("color"), fallback="blue"))
        add_text(slide, x + 0.25, y, 0.9, ch, item["icon"], font_size=32,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + 1.2, y + 0.18, cw - 1.4, 0.4, item["title"],
                 font_size=15, bold=True, color=theme.c("dark"), anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + 1.2, y + 0.6, cw - 1.4, 0.8, item["desc"],
                 font_size=11, color=theme.c("gray"), line_spacing=1.2)


def build_steps(prs, spec, ctx):
    theme = ctx.theme
    slide = content_slide(prs, spec, ctx)
    steps = spec["steps"]
    step_w, gap = 2.7, 0.18
    total_used = step_w * len(steps) + gap * (len(steps) - 1)
    start_x = (PAGE_W - total_used) / 2
    y, h = 1.95, 4.7
    for i, step in enumerate(steps):
        x = start_x + (step_w + gap) * i
        color = theme.c(step.get("color"), fallback="blue")
        add_round_rect(slide, x, y, step_w, h, theme.c("white"),
                       line_color=CARD_BORDER, corner=0.04)
        add_rect(slide, x, y, step_w, 1.1, color)
        nc = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + step_w / 2 - 0.5),
                                    Inches(y + 0.25), Inches(1.0), Inches(1.0))
        nc.fill.solid()
        nc.fill.fore_color.rgb = theme.c("white")
        nc.line.color.rgb = color
        nc.line.width = Pt(3)
        add_text(slide, x + step_w / 2 - 0.5, y + 0.25, 1.0, 1.0, step["no"],
                 font_size=36, bold=True, color=color, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + 0.15, y + 1.3, step_w - 0.3, 0.55, step["title"],
                 font_size=15, bold=True, color=theme.c("dark"),
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + 0.15, y + 1.95, step_w - 0.3, h - 2.3, step["desc"],
                 font_size=12, color=theme.c("gray"), align=PP_ALIGN.CENTER,
                 line_spacing=1.4)
        if step.get("url"):
            add_text(slide, x + 0.1, y + h - 0.55, step_w - 0.2, 0.4, step["url"],
                     font_size=10, color=theme.c("blue_light"),
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def build_cards_flow(prs, spec, ctx):
    theme = ctx.theme
    slide = content_slide(prs, spec, ctx)
    top_cards = spec["top_cards"]
    widths = spec.get("widths", [6.0, 5.9])
    card_h = top_cards[0].get("height", 2.0)
    render_card(slide, top_cards[0], 0.6, 1.85, widths[0], card_h, theme)
    render_card(slide, top_cards[1], PAGE_W - 0.6 - widths[1], 1.85, widths[1], card_h, theme)
    flow = spec["flow"]
    add_text(slide, 0.6, 4.05, 12, 0.4, flow["title"], font_size=16, bold=True,
             color=theme.c("blue"))
    steps = flow["steps"]
    flow_w, flow_gap = 2.0, 0.06
    flow_total = flow_w * len(steps) + flow_gap * (len(steps) - 1)
    flow_start = (PAGE_W - flow_total) / 2
    fy = 4.6
    for i, step in enumerate(steps):
        fx = flow_start + (flow_w + flow_gap) * i
        add_round_rect(slide, fx, fy, flow_w, 2.0, theme.c("white"),
                       line_color=CARD_BORDER, corner=0.1)
        add_rect(slide, fx, fy, flow_w, 0.45, theme.c("blue_light"))
        add_text(slide, fx, fy, flow_w, 0.45, f"Step {i + 1}", font_size=11,
                 bold=True, color=theme.c("white"), align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, fx + 0.1, fy + 0.5, flow_w - 0.2, 0.45, step["title"],
                 font_size=13, bold=True, color=theme.c("dark"),
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_code_block(slide, fx + 0.1, fy + 1.0, flow_w - 0.2, 0.95,
                       step["code"], font_size=10, label=f"step{i + 1}")
        if i < len(steps) - 1:
            arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                         Inches(fx + flow_w + 0.005), Inches(fy + 0.9),
                                         Inches(0.05), Inches(0.2))
            arr.fill.solid()
            arr.fill.fore_color.rgb = theme.c("blue_light")
            arr.line.fill.background()


def build_exercise_grid(prs, spec, ctx):
    theme = ctx.theme
    slide = content_slide(prs, spec, ctx)
    items = spec["items"]
    cw, ch = 6.05, 2.55
    gx, gy = 0.15, 0.15
    start_x, start_y = 0.6, 1.85
    for i, item in enumerate(items):
        x = start_x + (i % 2) * (cw + gx)
        y = start_y + (i // 2) * (ch + gy)
        color = theme.c(item.get("color"), fallback="blue")
        add_round_rect(slide, x, y, cw, ch, theme.c("white"),
                       line_color=CARD_BORDER, corner=0.04)
        add_rect(slide, x, y, 0.12, ch, color)
        add_round_rect(slide, x + 0.3, y + 0.2, 0.85, 0.45, color, corner=0.25)
        add_text(slide, x + 0.3, y + 0.2, 0.85, 0.45, item["label"], font_size=12,
                 bold=True, color=theme.c("white"), align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + 1.3, y + 0.18, cw - 2.5, 0.5, item["title"],
                 font_size=16, bold=True, color=theme.c("dark"), anchor=MSO_ANCHOR.MIDDLE)
        add_round_rect(slide, x + cw - 1.1, y + 0.25, 0.95, 0.35,
                       theme.c("bg_light"), corner=0.3)
        add_text(slide, x + cw - 1.1, y + 0.25, 0.95, 0.35, f"⏱  {item['time']}",
                 font_size=10, bold=True, color=theme.c("blue"),
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + 0.3, y + 0.85, cw - 0.6, ch - 1.0, item["desc"],
                 font_size=12, color=theme.c("dark"), line_spacing=1.45)


def build_summary(prs, spec, ctx):
    theme = ctx.theme
    slide = content_slide(prs, spec, ctx)
    items = spec["items"]
    cw, ch, gap = 2.4, 2.8, 0.12
    total_w = cw * len(items) + gap * (len(items) - 1)
    start_x = (PAGE_W - total_w) / 2
    y = 1.95
    for i, item in enumerate(items):
        x = start_x + (cw + gap) * i
        color = theme.c(item.get("color"), fallback="blue")
        add_round_rect(slide, x, y, cw, ch, theme.c("white"),
                       line_color=CARD_BORDER, corner=0.06)
        add_rect(slide, x, y, cw, 0.12, color)
        nc = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + cw / 2 - 0.45),
                                    Inches(y + 0.3), Inches(0.9), Inches(0.9))
        nc.fill.solid()
        nc.fill.fore_color.rgb = color
        nc.line.fill.background()
        add_text(slide, x + cw / 2 - 0.45, y + 0.3, 0.9, 0.9, item["no"],
                 font_size=24, bold=True, color=theme.c("white"),
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + 0.1, y + 1.3, cw - 0.2, 0.5, item["title"],
                 font_size=16, bold=True, color=theme.c("dark"),
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, x + 0.2, y + 1.85, cw - 0.4, 0.9, item["desc"],
                 font_size=11, color=theme.c("gray"), align=PP_ALIGN.CENTER,
                 line_spacing=1.3)
    banner = spec["banner"]
    by = 5.1
    add_round_rect(slide, 0.6, by, 12.15, 1.5, theme.c("blue"), corner=0.04)
    add_text(slide, 0.6, by + 0.2, 12.15, 0.5, banner["label"], font_size=14,
             bold=True, color=theme.c("yellow"), align=PP_ALIGN.CENTER)
    add_text(slide, 0.6, by + 0.6, 12.15, 0.8, "\n".join(banner["lines"]),
             font_size=16, color=theme.c("white"), align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.4)


def build_qa(prs, spec, ctx):
    theme = ctx.theme
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, PAGE_W, PAGE_H, theme.c("white"))
    add_rect(slide, 0, 0, 6.5, PAGE_H, theme.c("blue"))
    for ox, oy, d, color in [(5.5, -1.0, 3.5, "blue_light"), (5.8, 5.5, 2.5, "yellow")]:
        deco = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(ox), Inches(oy),
                                      Inches(d), Inches(d))
        deco.fill.solid()
        deco.fill.fore_color.rgb = theme.c(color)
        deco.line.fill.background()
    add_text(slide, 0.7, 2.0, 6, 1.5, spec["headline"], font_size=110,
             bold=True, color=theme.c("white"))
    add_rect(slide, 0.85, 3.85, 1.5, 0.08, theme.c("yellow"))
    add_text(slide, 0.7, 4.0, 6, 0.6, spec["sub_cn"], font_size=24,
             color=theme.c("yellow"))
    add_text(slide, 0.7, 4.7, 6, 0.5, spec["sub_en"], font_size=14,
             color=COVER_SUBTLE)
    add_text(slide, 0.7, 6.4, 6, 0.4, spec["thanks"], font_size=16, bold=True,
             color=theme.c("white"))
    quote = spec["quote"]
    add_text(slide, 7.0, 1.8, 5.8, 0.5, quote["label"], font_size=14, bold=True,
             color=theme.c("yellow"))
    add_text(slide, 7.0, 2.3, 5.8, 2.0, quote["en"], font_size=30, bold=True,
             color=theme.c("blue"))
    add_text(slide, 7.0, 3.4, 5.8, 1.0, quote["cn"], font_size=22,
             color=theme.c("dark"))
    add_rect(slide, 7.0, 4.6, 1.0, 0.06, theme.c("yellow"))
    add_text(slide, 7.0, 4.8, 5.8, 0.5, quote["attribution"], font_size=14,
             color=theme.c("gray"))
    resources = spec["resources"]
    add_text(slide, 7.0, 5.6, 5.8, 0.4, resources["label"], font_size=14,
             bold=True, color=theme.c("blue"))
    add_text(slide, 7.0, 6.0, 5.8, 0.4, resources["text"], font_size=12,
             color=theme.c("dark"))


# ============== 注册表 / 校验 / CLI ==============

BUILDERS = {
    "cover": build_cover,
    "divider": build_divider,
    "agenda": build_agenda,
    "two_cards": build_two_cards,
    "story_person": build_story_person,
    "table": build_table,
    "stats": build_stats,
    "ai_prompts": build_ai_prompts,
    "code_compare": build_code_compare,
    "code_demo": build_code_demo,
    "code_full": build_code_full,
    "icon_grid": build_icon_grid,
    "steps": build_steps,
    "cards_flow": build_cards_flow,
    "exercise_grid": build_exercise_grid,
    "summary": build_summary,
    "qa": build_qa,
}

REQUIRED_FIELDS = {
    "cover": ["kicker", "title_lines", "subtitle", "credit"],
    "divider": ["no", "title", "subtitle", "color"],
    "agenda": ["items", "sidebar"],
    "two_cards": ["cards"],
    "story_person": ["story", "person"],
    "table": ["headers", "rows", "widths"],
    "stats": ["cards", "banner"],
    "ai_prompts": ["prompts"],
    "code_compare": ["cards", "notes"],
    "code_demo": ["scenario", "code"],
    "code_full": ["code"],
    "icon_grid": ["items"],
    "steps": ["steps"],
    "cards_flow": ["top_cards", "flow"],
    "exercise_grid": ["items"],
    "summary": ["items", "banner"],
    "qa": ["headline", "sub_cn", "sub_en", "thanks", "quote", "resources"],
}


def validate_slide(i, spec):
    def err(msg):
        raise ValueError(f"slide {i + 1} (type {spec.get('type', '?')}): {msg}")
    stype = spec.get("type")
    if stype not in BUILDERS:
        err(f"未知版式类型 '{stype}'，合法类型: {sorted(BUILDERS)}")
    for field in REQUIRED_FIELDS[stype]:
        if field not in spec:
            err(f"缺少必填字段 '{field}'")


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    json_path = os.path.abspath(sys.argv[1])
    with open(json_path, encoding="utf-8-sig") as f:
        data = json.load(f)

    theme = Theme(data.get("palette"))
    course_line = f"{data['course']}  ·  {data['week']}"
    ctx = Ctx(page_no=0, total=len(data["slides"]), course_line=course_line, theme=theme)

    if len(sys.argv) > 2:
        out_path = os.path.abspath(sys.argv[2])
    elif data.get("output"):
        out_path = os.path.join(os.path.dirname(json_path), data["output"])
    else:
        out_path = os.path.splitext(json_path)[0] + ".pptx"

    prs = Presentation()
    prs.slide_width = Inches(PAGE_W)
    prs.slide_height = Inches(PAGE_H)

    for i, spec in enumerate(data["slides"]):
        try:
            validate_slide(i, spec)
            BUILDERS[spec["type"]](prs, spec, replace(ctx, page_no=i + 1))
        except ValueError as e:
            print(f"[ERROR] {e}")
            sys.exit(1)

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    prs.save(out_path)
    print(f"[OK] PPT generated successfully! Total slides: {len(prs.slides)}")
    print(f"[PATH] {out_path}")


if __name__ == "__main__":
    main()
